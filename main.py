from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import HTMLResponse, Response
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import httpx, io, os, base64

app = FastAPI(title="JARVIS Scholar", description="Cloud document intelligence by JARVIS Scholar.", version="9.0.0")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash").strip()
GEMINI_FALLBACK_MODEL = os.getenv("GEMINI_FALLBACK_MODEL", "gemini-3.5-flash-lite").strip()
uploaded_document_text = ""
uploaded_document_name = ""
uploaded_document_type = ""
chat_history = []

HTML = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>JARVIS Scholar</title>
<style>
:root{--bg:#07111d;--panel:#0d1b2a;--border:#1f3b52;--text:#f3f7fb;--muted:#98adbf;--accent:#57c7ff;--success:#62d8a0}
*{box-sizing:border-box}body{margin:0;min-height:100vh;background:#07111d;color:var(--text);font-family:system-ui,sans-serif}
.shell{width:min(1180px,calc(100% - 30px));margin:auto;padding:28px 0}.top{display:flex;justify-content:space-between;align-items:center;margin-bottom:22px}
.brand{display:flex;gap:13px;align-items:center}.orb{width:44px;height:44px;border-radius:50%;background:radial-gradient(circle at 38% 34%,#e1f8ff,#5ad2ff 15%,#087db8 45%,#031d2f 75%);box-shadow:0 0 24px #50caff66}
h1{margin:0;font-size:22px;letter-spacing:.15em}.muted,.status{color:var(--muted);font-size:13px}.dot{display:inline-block;width:9px;height:9px;border-radius:50%;background:#667583;margin-right:7px}.dot.online{background:var(--success)}
.layout{display:grid;grid-template-columns:330px 1fr;gap:20px}.card{background:var(--panel);border:1px solid var(--border);border-radius:18px}.side{padding:20px;min-height:680px}
.drop{margin-top:18px;padding:22px 14px;text-align:center;border:1.5px dashed #31566f;border-radius:14px;background:#0a1724}.drop strong{display:block;margin-bottom:8px}input[type=file]{width:100%;margin-top:14px;color:var(--muted)}
button{border:0;border-radius:11px;padding:11px 15px;font:inherit;font-weight:700;cursor:pointer}.primary{background:#57c7ff;color:#03131e}.secondary{background:#14283a;color:var(--text);border:1px solid var(--border)}.full{width:100%;margin-top:12px}
.info{margin-top:14px;padding:12px;border:1px solid #17334a;border-radius:11px;background:#091724;color:var(--muted);font-size:13px;white-space:pre-wrap}.quick{display:grid;grid-template-columns:1fr 1fr;gap:8px;margin-top:18px}.quick button{background:#10263a;color:#cdefff;border:1px solid #224761;font-size:12px}.langbox{margin-top:14px}.langbox label{display:block;margin-bottom:6px;color:var(--muted);font-size:12px}.langbox select{width:100%;background:#091724;color:var(--text);border:1px solid #224761;border-radius:11px;padding:10px 12px;font:inherit}
.chat{display:flex;flex-direction:column;min-height:680px;overflow:hidden}.head{padding:18px 20px;border-bottom:1px solid var(--border);display:flex;justify-content:space-between}.messages{flex:1;padding:22px;overflow:auto;max-height:540px}.empty{text-align:center;color:var(--muted);padding:90px 20px}
.msg{max-width:82%;margin-bottom:14px;padding:13px 15px;border-radius:14px;white-space:pre-wrap;line-height:1.6}.user{margin-left:auto;background:#15354b}.assistant{margin-right:auto;background:#0c2232}.role{display:block;font-size:11px;color:var(--muted);margin-bottom:5px}
.composer{padding:16px;border-top:1px solid var(--border);background:#0a1724}textarea{width:100%;min-height:84px;border-radius:13px;border:1px solid var(--border);background:#07131f;color:var(--text);padding:13px;font:inherit}.row{display:flex;justify-content:space-between;align-items:center;margin-top:10px;gap:10px}
@media(max-width:850px){.layout{grid-template-columns:1fr}.side{min-height:auto}}
</style></head><body><main class="shell">
<div class="top"><div class="brand"><div class="orb"></div><div><h1>JARVIS</h1><div class="muted">Scholar • Document Intelligence</div></div></div><div class="status"><span id="statusDot" class="dot"></span><span id="statusText">Checking JARVIS...</span></div></div>
<div class="layout"><aside class="card side"><h2>Document</h2><p class="muted">Chat with JARVIS straight away, or upload a document for document-aware assistance.</p>
<div class="drop"><strong>Optional document</strong><span class="muted">PDF • Word • PowerPoint • TXT</span><input id="documentFile" type="file" accept=".pdf,.docx,.pptx,.txt"></div>
<button id="uploadBtn" class="primary full" onclick="uploadDocument()">Upload document</button><button class="secondary full" onclick="removeDocument()">Remove document</button><div id="docInfo" class="info">No document • General chat mode</div>
<div class="langbox"><label for="language">Response & voice language</label><select id="language"><option value="auto">Auto-detect</option><option value="en">English</option><option value="ar">Arabic</option><option value="fr">French</option><option value="es">Spanish</option><option value="mr">Marathi</option></select></div>
<div class="quick"><button onclick="quickAsk('Summarise this document clearly.')">Summarise</button><button onclick="quickAsk('Explain the main points in simple terms.')">Explain simply</button><button onclick="quickAsk('What are the most important points in this document?')">Key points</button><button onclick="quickAsk('Create a concise study guide from this document.')">Study guide</button><button onclick="quickAsk('Critique this document and identify its weaknesses, gaps and risks.')">Critique</button><button onclick="quickAsk('How can I improve this document? Give specific actionable changes.')">Improve</button></div></aside>
<section class="card chat"><div class="head"><div><h2>Chat with JARVIS</h2><div class="muted">General AI chat • automatically becomes document-aware when you upload a file.</div></div><button class="secondary" onclick="clearChat()">Clear chat</button></div>
<div id="messages" class="messages"><div class="empty" id="emptyState">Ask JARVIS anything. Uploading a document is optional.</div></div>
<div style="padding:10px 16px;border-top:1px solid var(--border);display:flex;justify-content:space-between;align-items:center;gap:10px">
<span id="voiceStatus" class="muted">Voice ready</span>
<div style="display:flex;gap:7px">
<button class="secondary" onclick="pauseSpeech()">⏸ Pause</button>
<button class="secondary" onclick="resumeSpeech()">▶ Resume</button>
<button class="secondary" onclick="stopSpeech()">⏹ Stop</button>
</div>
</div>
<div class="composer"><textarea id="question" placeholder="Ask anything: improve it, rewrite it, explain it, translate it, compare it..."></textarea><div class="row"><span class="muted">JARVIS can chat normally, or analyse, critique, improve, rewrite, translate and teach from an uploaded document.</span><button id="askBtn" class="primary" onclick="askJarvis()">Ask JARVIS</button></div></div></section></div>
<footer style="text-align:center;color:#98adbf;font-size:12px;padding:24px 10px 4px">
<strong style="color:#f3f7fb">JARVIS Scholar</strong> — Built by <strong style="color:#cdefff">Saleh Pour</strong><br>
<span>Robotics &amp; Artificial Intelligence</span>
</footer></main>
<script>
let currentAudio=null;
let currentAudioUrl=null;

const LANGUAGE_SETTINGS={
 auto:{name:"Auto",locale:"en-GB"},
 en:{name:"English",locale:"en-GB"},
 ar:{name:"Arabic",locale:"ar-XA"},
 fr:{name:"French",locale:"fr-FR"},
 es:{name:"Spanish",locale:"es-ES"},
 mr:{name:"Marathi",locale:"mr-IN"}
};

function selectedLanguage(){
 return document.getElementById("language")?.value || "auto";
}

function detectLanguageFromText(text){
 const chosen=selectedLanguage();
 if(chosen!=="auto") return chosen;
 if(/[\u0600-\u06FF]/.test(text)) return "ar";
 if(/[\u0900-\u097F]/.test(text)) return "mr";
 const lower=" "+text.toLowerCase()+" ";
 const spanish=[" el "," la "," los "," las "," que "," para "," con "," una "," por "," como "," esto "," esta "," gracias "," mejorar "," documento "];
 const french=[" le "," les "," des "," une "," pour "," avec "," dans "," est "," comment "," merci "," améliorer "," document "," résumé "];
 let es=spanish.filter(w=>lower.includes(w)).length;
 let fr=french.filter(w=>lower.includes(w)).length;
 if(/[ñ¿¡]/i.test(text)) es+=3;
 if(/[àâçéèêëîïôùûüÿœ]/i.test(text)) fr+=2;
 if(es>fr && es>=2) return "es";
 if(fr>es && fr>=2) return "fr";
 return "en";
}

function cleanupAudio(){
 if(currentAudio){
   try{currentAudio.pause();}catch{}
   currentAudio=null;
 }
 if(currentAudioUrl){
   URL.revokeObjectURL(currentAudioUrl);
   currentAudioUrl=null;
 }
}

async function speakText(text){
 const status=document.getElementById("voiceStatus");
 if(!text?.trim()) return;
 cleanupAudio();
 status.textContent="Generating natural voice...";
 const lang=detectLanguageFromText(text);

 const fd=new FormData();
 fd.append("text",text);
 fd.append("language",lang);

 try{
   const r=await fetch("/tts",{method:"POST",body:fd});
   if(!r.ok){
     let msg="Voice generation failed.";
     try{const x=await r.json();msg=x.error||msg;}catch{}
     status.textContent=msg;
     return;
   }
   const blob=await r.blob();
   currentAudioUrl=URL.createObjectURL(blob);
   currentAudio=new Audio(currentAudioUrl);
   const cfg=LANGUAGE_SETTINGS[lang]||LANGUAGE_SETTINGS.en;
   currentAudio.onplay=()=>status.textContent=`JARVIS is speaking • ${cfg.name}`;
   currentAudio.onended=()=>{status.textContent="Voice ready";cleanupAudio();};
   currentAudio.onerror=()=>{status.textContent="Voice playback failed";cleanupAudio();};
   await currentAudio.play();
 }catch(e){
   status.textContent="Voice failed: "+e.message;
   cleanupAudio();
 }
}

function pauseSpeech(){
 if(currentAudio && !currentAudio.paused){
   currentAudio.pause();
   document.getElementById("voiceStatus").textContent="Voice paused";
 }
}

function resumeSpeech(){
 if(currentAudio && currentAudio.paused){
   currentAudio.play();
   document.getElementById("voiceStatus").textContent="JARVIS is speaking...";
 }
}

function stopSpeech(){
 cleanupAudio();
 document.getElementById("voiceStatus").textContent="Voice ready";
}

async function checkJarvis(){let d=document.getElementById("statusDot"),t=document.getElementById("statusText");try{let r=await fetch("/jarvis-status"),x=await r.json();if(x.connected){d.classList.add("online");t.textContent="JARVIS online • "+(x.jarvis?.brain||"Local AI")}else{d.classList.remove("online");t.textContent="JARVIS offline"}}catch{d.classList.remove("online");t.textContent="JARVIS offline"}}
function addMessage(role,text){let m=document.getElementById("messages"),e=document.getElementById("emptyState");if(e)e.remove();let d=document.createElement("div");d.className="msg "+role;let l=document.createElement("span");l.className="role";l.textContent=role==="user"?"You":"JARVIS";let c=document.createElement("div");c.textContent=text;d.append(l,c);
if(role==="assistant"){
 let tools=document.createElement("div");
 tools.style.marginTop="10px";
 let read=document.createElement("button");
 read.textContent="🔊 Read aloud";
 read.className="secondary";
 read.style.padding="6px 9px";
 read.style.fontSize="11px";
 read.onclick=()=>speakText(text);
 tools.appendChild(read);
 d.appendChild(tools);
}
m.appendChild(d);m.scrollTop=m.scrollHeight}
async function uploadDocument(){let f=document.getElementById("documentFile"),i=document.getElementById("docInfo"),b=document.getElementById("uploadBtn");if(!f.files.length){i.textContent="Choose a document first.";return}let fd=new FormData();fd.append("file",f.files[0]);b.disabled=true;i.textContent="Uploading and extracting text...";try{let r=await fetch("/upload",{method:"POST",body:fd}),x=await r.json();i.textContent=x.error?"Error: "+x.error:`Ready: ${x.filename}\n${x.document_type} • ${x.structure} • ${x.characters_extracted.toLocaleString()} characters`;if(!x.error)await clearChat(false)}catch(e){i.textContent="Upload failed: "+e.message}finally{b.disabled=false}}
async function quickAsk(q){document.getElementById("question").value=q;await askJarvis()}
async function askJarvis(){let q=document.getElementById("question"),text=q.value.trim(),b=document.getElementById("askBtn");if(!text)return;addMessage("user",text);q.value="";b.disabled=true;let fd=new FormData();fd.append("question",text);fd.append("language",selectedLanguage());try{let r=await fetch("/ask",{method:"POST",body:fd}),x=await r.json();addMessage("assistant",x.error?"Error: "+x.error:(x.answer||"JARVIS returned an empty response."))}catch(e){addMessage("assistant","Request failed: "+e.message)}finally{b.disabled=false}}
async function removeDocument(){
 try{
   await fetch("/remove-document",{method:"POST"});
   document.getElementById("documentFile").value="";
   document.getElementById("docInfo").textContent="No document • General chat mode";
   await clearChat(false);
 }catch(e){
   document.getElementById("docInfo").textContent="Could not remove document: "+e.message;
 }
}
async function clearChat(clear=true){try{await fetch("/clear-chat",{method:"POST"})}catch{}document.getElementById("messages").innerHTML='<div class="empty" id="emptyState">Ask JARVIS anything. Uploading a document is optional.</div>';if(clear)document.getElementById("question").value=""}
document.getElementById("question").addEventListener("keydown",e=>{if(e.key==="Enter"&&!e.shiftKey){e.preventDefault();askJarvis()}});checkJarvis();setInterval(checkJarvis,10000);
</script></body></html>"""

def extract_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for n, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"\n--- PAGE {n} ---\n{text}")
    return "\n".join(parts), f"{len(reader.pages)} pages"

def extract_docx(data):
    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), f"{len(doc.paragraphs)} paragraphs"

def extract_pptx(data):
    pres = Presentation(io.BytesIO(data))
    parts = []
    for n, slide in enumerate(pres.slides, 1):
        texts = [s.text.strip() for s in slide.shapes if hasattr(s, "text") and s.text.strip()]
        if texts:
            parts.append(f"\n--- SLIDE {n} ---\n" + "\n".join(texts))
    return "\n".join(parts), f"{len(pres.slides)} slides"

def extract_txt(data):
    text = data.decode("utf-8", errors="replace")
    return text, f"{len(text.splitlines())} lines"

@app.get("/", response_class=HTMLResponse)
async def website():
    return HTMLResponse(HTML)

@app.get("/jarvis-status")
async def jarvis_status():
    if not GEMINI_API_KEY:
        return {"connected": False, "error": "GEMINI_API_KEY is not configured."}
    return {"connected": True, "jarvis": {"brain": f"Cloud AI • {GEMINI_MODEL}", "fallback": GEMINI_FALLBACK_MODEL}}

@app.post("/upload")
async def upload_document(file: UploadFile = File(...)):
    global uploaded_document_text, uploaded_document_name, uploaded_document_type, chat_history
    if not file.filename:
        return {"error": "Please choose a document."}
    ext = os.path.splitext(file.filename.lower())[1]
    if ext not in {".pdf", ".docx", ".pptx", ".txt"}:
        return {"error": "Unsupported file type. Please upload PDF, DOCX, PPTX or TXT."}
    data = await file.read()
    if not data:
        return {"error": "The uploaded document is empty."}
    try:
        if ext == ".pdf":
            text, structure, kind = *extract_pdf(data), "PDF"
        elif ext == ".docx":
            text, structure, kind = *extract_docx(data), "Word document"
        elif ext == ".pptx":
            text, structure, kind = *extract_pptx(data), "PowerPoint"
        else:
            text, structure, kind = *extract_txt(data), "Text document"
    except Exception as error:
        return {"error": f"The document could not be read: {error}"}
    if not text.strip():
        return {"error": "No readable text could be extracted from this document."}
    uploaded_document_text, uploaded_document_name, uploaded_document_type = text, file.filename, kind
    chat_history = []
    return {"filename": file.filename, "document_type": kind, "structure": structure, "characters_extracted": len(text)}


@app.post("/tts")
async def text_to_speech(text: str = Form(...), language: str = Form("auto")):
    if not GEMINI_API_KEY:
        return Response(
            content='{"error":"Cloud voice is not configured."}',
            status_code=503,
            media_type="application/json",
        )

    spoken_text = text.strip()
    if not spoken_text:
        return Response(
            content='{"error":"There is no text to read."}',
            status_code=400,
            media_type="application/json",
        )

    spoken_text = spoken_text[:9000]

    voice_profiles = {
        "en": {
            "voice": "Enceladus",
            "direction": """A very human adult British male voice inspired by a refined personal AI butler: polished, posh, intelligent, calm and conversational.
Use natural British English pronunciation with a subtle RP / educated London character, never American.

The performance must feel like a person speaking spontaneously to one listener, not like a narrator reading text.
Use realistic phrasing, uneven micro-pauses, tiny timing imperfections, gentle changes in tempo, and subtle emphasis shifts.
At the ends of some sentences, let the voice soften naturally instead of landing every line with the same cadence.
Between selected sentences or thought groups, include a quiet natural inhale or breath where a real speaker would need one.
The breathing should be audible but subtle: occasional, irregular and organic, never forced, never after every sentence, and never spoken as literal words.
Where appropriate, allow very light conversational hesitation or thinking rhythm before a complex point, but do not insert distracting filler into factual answers.
Short answers should sound effortless. Longer answers should breathe, pace themselves and feel physically spoken.
Avoid announcer cadence, audiobook cadence, commercial polish, exaggerated acting, robotic timing, over-enunciation or sing-song prosody.
The overall result should sound like an intelligent human assistant standing nearby and speaking naturally."""
        },
        "ar": {
            "voice": "Iapetus",
            "direction": "Natural fluent Modern Standard Arabic, authentic Arabic pronunciation, calm intelligent male-assistant delivery, natural pacing."
        },
        "fr": {
            "voice": "Iapetus",
            "direction": "Native metropolitan French pronunciation with a convincing French accent, relaxed intelligent male delivery and natural pacing."
        },
        "es": {
            "voice": "Iapetus",
            "direction": "Native European Spanish pronunciation with a convincing Spain Spanish accent, relaxed intelligent male delivery and natural pacing."
        },
        "mr": {
            "voice": "Iapetus",
            "direction": "Native Marathi pronunciation as spoken naturally in Maharashtra, India, with fluent rhythm and a calm intelligent male delivery."
        },
        "auto": {
            "voice": "Iapetus",
            "direction": "Detect the language of the text and use a convincing native pronunciation and natural human conversational delivery."
        },
    }
    profile = voice_profiles.get(language, voice_profiles["auto"])

    tts_prompt = f"""AUDIO PROFILE
You are the speaking voice of JARVIS Scholar.

DIRECTOR'S NOTES
{profile["direction"]}

PERFORMANCE RULES
- Preserve the meaning and wording of the supplied answer.
- Do not translate, summarise, explain, add an introduction, or omit substantive content.
- Do not read markdown symbols aloud.
- Convert punctuation and paragraph structure into natural spoken phrasing instead of mechanically reading punctuation.
- For English only, use subtle nonverbal performance such as a quiet inhale, soft exhale, tiny hesitation or brief thinking pause when it improves realism.
- Never announce or verbalise stage directions such as "breath", "inhale", "pause" or "sigh".
- Do not overuse nonverbal sounds; realism comes from restraint and irregularity.
- Keep lists intelligible with short natural pauses.
- The performance should sound like one person naturally speaking to the user.

TEXT TO SPEAK
{spoken_text}
"""

    payload = {
        "model": "gemini-3.1-flash-tts-preview",
        "input": tts_prompt,
        "response_format": {"type": "audio"},
        "generation_config": {
            "speech_config": [
                {"voice": profile["voice"]}
            ]
        },
    }

    try:
        async with httpx.AsyncClient(timeout=120.0) as client:
            response = await client.post(
                "https://generativelanguage.googleapis.com/v1beta/interactions",
                headers={
                    "x-goog-api-key": GEMINI_API_KEY,
                    "Content-Type": "application/json",
                    "Api-Revision": "2026-05-20",
                },
                json=payload,
            )

        if response.status_code != 200:
            try:
                detail = response.json().get("error", {}).get("message", "")
            except Exception:
                detail = response.text[:300]
            return Response(
                content='{"error":' + repr(f"Cloud voice returned HTTP {response.status_code}: {detail or 'request failed'}").replace("'", '"') + "}",
                status_code=502,
                media_type="application/json",
            )

        result = response.json()
        audio_block = None
        for step in result.get("steps", []):
            for block in step.get("content", []):
                if block.get("type") == "audio" and block.get("data"):
                    audio_block = block
                    break
            if audio_block:
                break

        if not audio_block:
            return Response(
                content='{"error":"Cloud voice returned no audio."}',
                status_code=502,
                media_type="application/json",
            )

        audio_bytes = base64.b64decode(audio_block["data"])
        mime_type = audio_block.get("mime_type") or "audio/wav"
        return Response(content=audio_bytes, media_type=mime_type)

    except Exception as error:
        return Response(
            content='{"error":' + repr(f"Voice request failed: {error}").replace("'", '"') + "}",
            status_code=502,
            media_type="application/json",
        )


@app.post("/ask")
async def ask_question(question: str = Form(...), language: str = Form("auto")):
    global chat_history
    if not question.strip():
        return {"error": "Please enter a question."}
    if not GEMINI_API_KEY:
        return {"error": "Cloud AI is not configured on the server."}

    supported_languages = {
        "auto": "Automatically detect the language of the user's latest message and answer in that same language.",
        "en": "Answer entirely in natural English.",
        "ar": "Answer entirely in natural Arabic.",
        "fr": "Answer entirely in natural French.",
        "es": "Answer entirely in natural Spanish.",
        "mr": "Answer entirely in natural Marathi (मराठी).",
    }
    language_instruction = supported_languages.get(language, supported_languages["auto"])

    history = "\n".join(
        f"{x['role'].upper()}: {x['content']}" for x in chat_history[-16:]
    )

    has_document = bool(uploaded_document_text.strip())
    context_truncated = False

    if has_document:
        context_limit = 120000
        document_context = uploaded_document_text[:context_limit]
        context_truncated = len(uploaded_document_text) > context_limit
        mode_instruction = f"""DOCUMENT-AWARE MODE

A document is currently uploaded.

DOCUMENT NAME:
{uploaded_document_name}

DOCUMENT CONTENT:
{document_context}

Use the uploaded document as the primary source for claims about that document.
You may still use general reasoning and knowledge when the user asks for critique, improvement, explanation, comparison, brainstorming or broader context.
Never invent quotations, numbers, names or document-specific facts that are not present.
Clearly distinguish your own analysis from information actually contained in the document."""
    else:
        mode_instruction = """GENERAL ASSISTANT MODE

No document is currently uploaded.
Act as a capable general-purpose AI assistant. The user does NOT need to upload a file before talking to you.
Answer questions, reason through problems, explain concepts, brainstorm, write, rewrite, plan, compare options and hold a natural multi-turn conversation.
Do not pretend a document exists.
If the user asks about current/live information you cannot verify from the conversation, say that live information may need to be checked rather than inventing it."""

    prompt = f"""You are JARVIS Scholar, an advanced AI assistant created by Saleh Pour.

PERSONALITY AND QUALITY
- Be intelligent, useful, natural and conversational.
- Reason carefully before answering.
- Give direct answers first, then enough explanation to be genuinely useful.
- Handle ambiguous follow-ups by using recent conversation context.
- Do not invent facts.
- Avoid generic filler and repetitive disclaimers.
- Match the user's requested level of detail.
- For writing or improvement requests, provide concrete rewritten examples when useful.
- For comparisons, explain trade-offs and make a recommendation when the user asks.
- For difficult problems, break the task into sensible steps without exposing hidden chain-of-thought.
- You can analyse, critique, improve, rewrite, summarise, translate, teach, brainstorm, compare, quiz, plan and answer general questions.

{mode_instruction}

LANGUAGE
{language_instruction}
If the user explicitly requests another supported language in the message, follow that request.
For Arabic, use natural Modern Standard Arabic unless a dialect is requested.
For Marathi, use natural Devanagari Marathi unless romanisation is requested.

RECENT CONVERSATION
{history or '(none)'}

LATEST USER MESSAGE
{question.strip()}
"""

    models_to_try = []
    for m in [GEMINI_MODEL, GEMINI_FALLBACK_MODEL]:
        if m and m not in models_to_try:
            models_to_try.append(m)

    last_error = ""
    for model_name in models_to_try:
        payload = {
            "contents": [{"role": "user", "parts": [{"text": prompt}]}],
            "generationConfig": {
                "temperature": 0.45,
                "topP": 0.92,
                "maxOutputTokens": 4096,
            },
        }

        try:
            async with httpx.AsyncClient(timeout=120.0) as client:
                response = await client.post(
                    f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent",
                    headers={
                        "x-goog-api-key": GEMINI_API_KEY,
                        "Content-Type": "application/json",
                    },
                    json=payload,
                )

            if response.status_code != 200:
                try:
                    last_error = response.json().get("error", {}).get("message", "")
                except Exception:
                    last_error = response.text[:300]
                continue

            result = response.json()
            candidates = result.get("candidates") or []
            if not candidates:
                last_error = "Cloud AI returned no answer."
                continue

            parts = candidates[0].get("content", {}).get("parts", [])
            answer = "\n".join(
                part.get("text", "") for part in parts if part.get("text")
            ).strip()

            if not answer:
                last_error = "Cloud AI returned an empty answer."
                continue

            chat_history += [
                {"role": "user", "content": question.strip()},
                {"role": "assistant", "content": answer},
            ]

            return {
                "document": uploaded_document_name if has_document else None,
                "document_type": uploaded_document_type if has_document else None,
                "answer": answer,
                "brain": model_name,
                "language": language,
                "mode": "document" if has_document else "general",
                "context_truncated": context_truncated,
            }

        except Exception as error:
            last_error = str(error)

    return {"error": f"Cloud AI request failed: {last_error or 'all configured models failed'}"}


@app.post("/clear-chat")
def clear_chat():
    global chat_history
    chat_history = []
    return {"message": "Chat history cleared."}

@app.post("/remove-document")
def remove_document():
    global uploaded_document_text, uploaded_document_name, uploaded_document_type, chat_history
    uploaded_document_text = ""
    uploaded_document_name = ""
    uploaded_document_type = ""
    chat_history = []
    return {"message": "Document removed. General chat mode active."}


@app.get("/document")
def document_info():
    if not uploaded_document_text:
        return {"message": "No document uploaded. General chat mode is active."}
    return {"filename": uploaded_document_name, "document_type": uploaded_document_type, "characters": len(uploaded_document_text), "chat_messages": len(chat_history)}
